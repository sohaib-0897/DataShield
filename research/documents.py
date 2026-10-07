"""Bounded offline document extraction and redacted rule evidence.

CERT attack labels never train a sensitivity model. No application policy is read
or changed. Worker text is private return data; public evidence contains no text.
"""
import argparse
import hashlib
import json
import logging
from pathlib import Path, PurePosixPath, PureWindowsPath
import re
import subprocess
import sys
import zipfile

VERSION = 'document-extraction-evidence-v1'
DEFAULTS = {'max_bytes': 2 * 1024**2, 'max_chars': 32768, 'max_pages': 40,
            'max_zip_bytes': 16 * 1024**2, 'max_zip_entries': 512,
            'cpu_seconds': 5, 'memory_bytes': 512 * 1024**2}
PATTERNS = {'cnic_like': re.compile(r'\b\d{5}-\d{7}-\d\b'),
            'email_address': re.compile(r'\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}\b')}


def private_result(status, text='', *, truncated=False, **extra):
    return {'status': status, 'text': text, 'truncated': truncated, **extra}


def worker_extract(path, limits):
    # Parse in an isolated process. POSIX limits are applied before heavy imports.
    import resource
    resource.setrlimit(resource.RLIMIT_CPU, (limits['cpu_seconds'], limits['cpu_seconds']))
    resource.setrlimit(resource.RLIMIT_AS, (limits['memory_bytes'], limits['memory_bytes']))
    logging.disable(logging.CRITICAL)
    if path.stat().st_size > limits['max_bytes']:
        return private_result('TOO_LARGE')
    suffix = path.suffix.lower()
    if suffix == '.txt':
        raw = path.read_bytes()
        if b'\0' in raw:
            return private_result('MALFORMED')
        text = raw.decode('utf-8-sig')
    elif suffix == '.pdf':
        from pypdf import PdfReader
        reader = PdfReader(path, strict=True)
        if reader.is_encrypted:
            return private_result('ENCRYPTED')
        if len(reader.pages) > limits['max_pages']:
            return private_result('LIMIT_EXCEEDED', limit='pages')
        text = ''
        for page in reader.pages:
            text += (page.extract_text() or '') + '\n'
            if len(text) > limits['max_chars']:
                return private_result('OK', text[:limits['max_chars']], truncated=True)
        if not text.strip():
            return private_result('SCANNED_OR_NO_TEXT')
    elif suffix == '.docx':
        with path.open('rb') as stream:
            signature = stream.read(8)
        if signature.startswith(b'\xd0\xcf\x11\xe0'):
            return private_result('ENCRYPTED_OR_LEGACY_UNSUPPORTED')
        with zipfile.ZipFile(path) as archive:
            members = archive.infolist()
            if len(members) > limits['max_zip_entries'] or sum(m.file_size for m in members) > limits['max_zip_bytes']:
                return private_result('LIMIT_EXCEEDED', limit='zip_expansion')
            names = set()
            for member in members:
                name = member.filename
                parts = PurePosixPath(name).parts
                if name.startswith('/') or '..' in parts or '\\' in name or ':' in name or name in names:
                    return private_result('MALFORMED')
                names.add(name)
                if ((member.external_attr >> 16) & 0o170000) == 0o120000:
                    return private_result('MALFORMED')
                if member.flag_bits & 1:
                    return private_result('ENCRYPTED')
                if member.file_size > 4 * 1024**2 or member.file_size > max(1, member.compress_size) * 200:
                    return private_result('LIMIT_EXCEEDED', limit='zip_member')
                if name.endswith('.xml'):
                    raw = archive.read(member)
                    if b'<!DOCTYPE' in raw.upper() or b'<!ENTITY' in raw.upper():
                        return private_result('MALFORMED')
        from docx import Document
        document = Document(path)
        text = ''
        paragraphs = (p.text for p in document.paragraphs)
        cells = (cell.text for table in document.tables for row in table.rows for cell in row.cells)
        from itertools import chain
        for part in chain(paragraphs, cells):
            text += part + '\n'
            if len(text) > limits['max_chars']:
                return private_result('OK', text[:limits['max_chars']], truncated=True)
    else:
        return private_result('UNSUPPORTED')
    return private_result('OK' if text.strip() else 'EMPTY', text[:limits['max_chars']],
                          truncated=len(text) > limits['max_chars'])


def extract_document(path, *, timeout=8, **overrides):
    """Return private text, explicit status and bounded extraction metadata."""
    path = Path(path).absolute()
    limits = {**DEFAULTS, **overrides}
    if set(limits) != set(DEFAULTS) or any(type(v) is not int or v < 1 for v in limits.values()) or not 0 < timeout <= 60:
        raise ValueError('Invalid extraction bounds')
    if any(p.is_symlink() for p in (path, *path.parents)):
        return private_result('UNSAFE_PATH')
    try:
        before = path.stat()
        if not path.is_file():
            return private_result('UNREADABLE')
    except OSError:
        return private_result('UNREADABLE')
    if before.st_size > limits['max_bytes']:
        return private_result('TOO_LARGE')
    if path.suffix.lower() not in {'.txt', '.pdf', '.docx'}:
        return private_result('UNSUPPORTED')
    try:
        result = subprocess.run([sys.executable, '-I', str(Path(__file__).resolve()), '--worker', str(path),
                                 '--limits', json.dumps(limits)], capture_output=True, timeout=timeout)
    except subprocess.TimeoutExpired:
        return private_result('TIMEOUT')
    if result.returncode != 0:
        return private_result('LIMIT_EXCEEDED' if result.returncode < 0 else 'MALFORMED')
    try:
        value = json.loads(result.stdout)
        after = path.stat()
    except (OSError, ValueError):
        return private_result('MALFORMED')
    if (before.st_size, before.st_mtime_ns, before.st_ino) != (after.st_size, after.st_mtime_ns, after.st_ino):
        return private_result('INPUT_CHANGED')
    if len(value.get('text', '')) > limits['max_chars']:
        return private_result('LIMIT_EXCEEDED')
    return value


def pii_evidence(text, *, max_spans=20):
    """No values/snippets, only rule types, counts and character offsets."""
    evidence = []
    for kind, pattern in PATTERNS.items():
        count, spans = 0, []
        for match in pattern.finditer(text):
            count += 1
            if len(spans) < max_spans:
                spans.append([match.start(), match.end()])
        if count:
            evidence.append({'type': kind, 'count': count, 'spans': spans, 'spans_truncated': count > max_spans})
    return evidence


def document_evidence(path, *, filename_labels=None, known_hashes=None, **bounds):
    """Additive independent filename/hash/content signals, no blocking decision."""
    path = Path(path)
    try:
        stat = path.stat()
        original_stat = (stat.st_size, stat.st_mtime_ns, stat.st_ino)
    except OSError:
        original_stat = None
    extraction = extract_document(path, **bounds)
    filename_labels, known_hashes = filename_labels or {}, known_hashes or {}
    name = PureWindowsPath(path.name).name.casefold()
    filename_label = next((label for filename, label in filename_labels.items()
                           if PureWindowsPath(filename).name.casefold() == name), None)
    digest = None
    # Keep filename evidence even for unreadable/unsupported/oversized documents.
    # A bounded full-file hash is independent of extracted content/PII findings.
    if extraction['status'] not in {'UNSAFE_PATH', 'UNREADABLE', 'INPUT_CHANGED'}:
        try:
            if path.stat().st_size <= 16 * 1024**2:
                hasher = hashlib.sha256()
                with path.open('rb') as stream:
                    for block in iter(lambda: stream.read(1024**2), b''):
                        hasher.update(block)
                digest = hasher.hexdigest()
        except OSError:
            pass
    if original_stat is not None and extraction['status'] not in {'UNSAFE_PATH', 'UNREADABLE'}:
        try:
            stat = path.stat()
            changed = original_stat != (stat.st_size, stat.st_mtime_ns, stat.st_ino)
        except OSError:
            changed = True
        if changed:
            extraction = private_result('INPUT_CHANGED')
            digest = None
    return {'version': VERSION, 'extraction_status': extraction['status'],
            'extracted_chars': len(extraction['text']), 'extraction_truncated': extraction['truncated'],
            'content_available': extraction['status'] == 'OK', 'pii_source': 'RULE_PATTERNS',
            'pii_evidence': pii_evidence(extraction['text']),
            'filename_label': filename_label, 'sha256': digest, 'hash_label': known_hashes.get(digest),
            'hash_available': digest is not None, 'sensitivity_model_status': 'UNAVAILABLE_NO_LABELED_CORPUS',
            'automated_blocking': False}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('path', nargs='?', type=Path)
    parser.add_argument('--worker', type=Path)
    parser.add_argument('--limits')
    args = parser.parse_args()
    if args.worker:
        try:
            value = worker_extract(args.worker, json.loads(args.limits))
        except ImportError:
            value = private_result('MISSING_DEPENDENCY')
        except MemoryError:
            value = private_result('LIMIT_EXCEEDED')
        except Exception:
            value = private_result('MALFORMED')
        sys.stdout.write(json.dumps(value))
    else:
        if args.path is None:
            parser.error('A document path is required')
        print(json.dumps(document_evidence(args.path), sort_keys=True))


if __name__ == '__main__':
    main()
