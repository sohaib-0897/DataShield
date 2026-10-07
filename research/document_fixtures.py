"""Synthetic test-only document fixtures. No genuine sensitivity supervision."""
from pathlib import Path

TEXT = 'Synthetic test report. Contact test@example.invalid. Identifier 12345-1234567-1.'


def write_pdf(path, *, text=TEXT, encrypted=False, pages=1):
    from pypdf import PdfWriter
    from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
    writer = PdfWriter()
    for _ in range(pages):
        page = writer.add_blank_page(width=612, height=792)
        if text:
            font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'),
                                     NameObject('/BaseFont'): NameObject('/Helvetica')})
            page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): writer._add_object(font)})})
            content = DecodedStreamObject()
            escaped = text.replace('\\', '\\\\').replace('(', '\\(').replace(')', '\\)')
            content.set_data(f'BT /F1 12 Tf 50 700 Td ({escaped}) Tj ET'.encode('ascii'))
            page[NameObject('/Contents')] = writer._add_object(content)
    if encrypted:
        writer.encrypt('synthetic-fixture-password')
    with Path(path).open('wb') as stream:
        writer.write(stream)


def write_docx(path, *, text=TEXT):
    from docx import Document
    document = Document()
    if text:
        document.add_paragraph(text)
        document.add_table(rows=1, cols=1).cell(0, 0).text = 'Synthetic table text'
    document.save(path)


def make_fixtures(folder):
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    (folder / 'report.txt').write_text(TEXT, encoding='utf-8')
    write_pdf(folder / 'report.pdf')
    write_docx(folder / 'report.docx')
    write_pdf(folder / 'encrypted.pdf', encrypted=True)
    write_pdf(folder / 'no_text.pdf', text='')
    (folder / 'empty.txt').write_text('')
    (folder / 'unsupported.xlsx').write_bytes(b'Synthetic unsupported document')
    (folder / 'malformed.pdf').write_bytes(b'not a pdf')
    (folder / 'malformed.docx').write_bytes(b'not a zip')
    (folder / 'malformed.txt').write_bytes(b'\xff\x00')
    return sorted(folder.iterdir())
